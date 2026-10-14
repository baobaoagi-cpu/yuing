"""Authenticated PWA preview with optional, isolated test accounts."""
from __future__ import annotations

from collections import defaultdict, deque
from datetime import timedelta
from functools import wraps
from hashlib import sha256
import hmac
import io
import json
import math
import os
import re
import sys
import threading
import time
import zipfile
from email.parser import BytesParser
from email import policy
from xml.etree import ElementTree as ET
from urllib.parse import urlparse
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError
from pathlib import Path
from flask import Flask, abort, jsonify, redirect, request, send_from_directory, session, send_file
from werkzeug.security import check_password_hash
from uuid import UUID

ROOT = Path(__file__).resolve().parent
PROJECT_ROOT = ROOT.parent
BACKEND_ROOT = PROJECT_ROOT / "backend"
sys.path.insert(0, str(BACKEND_ROOT))
from routes import yuling_bp  # noqa: E402
from deploy_pwa.chat_store import ChatStore, StoreUnavailable, PageReadQuotaExceeded
from deploy_pwa.progressive_memory import memories_for_turn
from deploy_pwa.chat_api import reply_to, ProviderError
from deploy_pwa.nico_profile_store import NicoProfileStore
from deploy_pwa.skill_growth import SkillGrowthStore, explicit_skill_need
from deploy_pwa.call_reflection import PROMPT_VERSION, substantial_call, generate_reflection
from deploy_pwa.public_page_reader import PageReadError, _validated_url, read_public_page
from deploy_pwa.time_tool import ZONES, TimeToolError, convert_time
USERNAME = os.environ.get("POC_PREVIEW_USER", "")
PASSWORD = os.environ.get("POC_PREVIEW_PASSWORD", "")
SESSION_SECRET = os.environ.get("POC_SESSION_SECRET", "")


def _load_extra_accounts():
    """Read optional username -> Werkzeug PBKDF2 hash entries, failing closed."""
    raw = os.environ.get("POC_EXTRA_ACCOUNTS_JSON", "").strip()
    if not raw:
        return {}
    try:
        accounts = json.loads(raw)
    except ValueError as exc:
        raise RuntimeError("POC_EXTRA_ACCOUNTS_JSON must be valid JSON") from exc
    if not isinstance(accounts, dict) or not 1 <= len(accounts) <= 20:
        raise RuntimeError("POC_EXTRA_ACCOUNTS_JSON must contain 1 to 20 accounts")
    folded_names = {USERNAME.casefold()} if USERNAME else set()
    for username, password_hash in accounts.items():
        if (not isinstance(username, str) or
                not re.fullmatch(r"[A-Za-z][A-Za-z0-9_.-]{2,63}", username) or
                username.casefold() in folded_names or not isinstance(password_hash, str) or
                not re.fullmatch(r"pbkdf2:sha256:1000000\$[A-Za-z0-9]{16}\$[0-9a-f]{64}",
                                 password_hash)):
            raise RuntimeError("POC_EXTRA_ACCOUNTS_JSON has an invalid account entry")
        folded_names.add(username.casefold())
    return accounts


EXTRA_ACCOUNTS = _load_extra_accounts()
BRIDGE = os.environ.get("POC_LOCAL_BRIDGE_URL", "").rstrip("/")
_bridge_url = urlparse(BRIDGE)
if BRIDGE and (_bridge_url.scheme != "http" or _bridge_url.hostname not in {"127.0.0.1", "localhost"}):
    raise RuntimeError("POC_LOCAL_BRIDGE_URL must be a loopback HTTP URL")
MAX_JSON = 16_384
MAX_AUDIO = 2_000_000
MAX_FILE = 5_000_000
REMEMBER_SECONDS = 400 * 24 * 60 * 60
SHORT_SECONDS = 12 * 60 * 60
app = Flask(__name__, static_folder=None)
app.register_blueprint(yuling_bp)
app.config.update(
    SECRET_KEY=SESSION_SECRET or os.urandom(32),
    SESSION_COOKIE_NAME="yuling_preview",
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SECURE=not bool(BRIDGE),
    SESSION_COOKIE_SAMESITE="Lax",
    PERMANENT_SESSION_LIFETIME=timedelta(days=400),
    SESSION_REFRESH_EACH_REQUEST=False,
    MAX_CONTENT_LENGTH=MAX_FILE + 10_000,
)
_failures = defaultdict(deque)
_failure_lock = threading.Lock()
_page_read_lock = threading.Lock()


@app.before_request
def _protect_original_yuling_api():
    if request.path.startswith("/api/yuling/"):
        # Agent-to-server usage receipts use their own HMAC and no browser cookie.
        if request.path == "/api/yuling/usage/ingest" and request.method == "POST":
            return None
        guard = _require_login(api=True)
        if guard:
            return guard
        if request.path == "/api/yuling/token" and request.method == "POST":
            if not _same_origin():
                return jsonify(error="来源不符"), 403
            request.environ["poc.same_origin"] = True
        # This value is written after the server validates its login cookie.
        # The token route must never derive memory ownership from browser JSON.
        request.environ["poc.authenticated_owner"] = session["account"]


def _configured():
    return bool(USERNAME and PASSWORD and SESSION_SECRET and len(SESSION_SECRET) >= 32)


def _credential_tag(account):
    if account == USERNAME:
        # Preserve cookies issued before optional test accounts were added.
        material = f"{USERNAME}\0{PASSWORD}"
    elif isinstance(account, str) and account in EXTRA_ACCOUNTS:
        material = f"extra\0{account}\0{EXTRA_ACCOUNTS[account]}"
    else:
        return None
    return hmac.new(SESSION_SECRET.encode(), material.encode(), sha256).hexdigest()[:24]


def _signed_in():
    account = session.get("account")
    tag = _credential_tag(account)
    return (_configured() and tag is not None and
            isinstance(session.get("credential_tag"), str) and
            hmac.compare_digest(session["credential_tag"], tag) and
            isinstance(session.get("until"), (int, float)) and session["until"] > time.time())


def _require_login(api=False):
    if _signed_in():
        return None
    if api:
        return jsonify(error="请先登录"), 401
    login_url = "/login?next=" + request.path if request.path in {"/call", "/chat"} else "/login"
    return redirect(login_url, code=303)


def _same_origin():
    origin = request.headers.get("Origin", "")
    scheme = "http" if BRIDGE else "https"
    return origin == f"{scheme}://{request.host}" and request.headers.get("Sec-Fetch-Site", "same-origin") != "cross-site"


@app.after_request
def _headers(response):
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Referrer-Policy"] = "same-origin"
    if (request.path.startswith("/api/") or response.mimetype == "text/html" or
            request.path in {"/chat.css", "/chat.js", "/swipe-navigation.js", "/update.js"}):
        response.headers["Cache-Control"] = "no-store"
    return response


@app.get("/healthz")
def healthz():
    return ("", 200 if _configured() else 503)


@app.get("/login")
def login_page():
    if not _configured():
        abort(503)
    if _signed_in():
        requested = request.args.get("next")
        return redirect(requested if requested in {"/call", "/chat"} else "/chat", code=303)
    return send_from_directory(ROOT, "login.html")


@app.get("/call")
def original_call_page():
    guard = _require_login()
    if guard:
        return guard
    return send_from_directory(PROJECT_ROOT / "frontend", "index.html")


@app.get("/sense")
def sound_preview_page():
    guard = _require_login()
    if guard:
        return guard
    return send_from_directory(ROOT, "index.html")


@app.get("/chat")
def conversation_page():
    guard = _require_login()
    if guard:
        return guard
    return send_from_directory(PROJECT_ROOT / "frontend", "chat.html")


@app.get("/tools")
def tool_page():
    guard = _require_login()
    if guard:
        return guard
    return send_from_directory(PROJECT_ROOT / "frontend", "tools.html")


@app.get("/tools.css")
@app.get("/tools.js")
@app.get("/tool-menu.js")
def tool_asset():
    guard = _require_login()
    if guard:
        return guard
    return send_from_directory(PROJECT_ROOT / "frontend", request.path[1:])


@app.get("/chat.css")
@app.get("/chat.js")
@app.get("/swipe-navigation.js")
@app.get("/update.js")
@app.get("/emotion-video.js")
def conversation_asset():
    guard = _require_login()
    if guard:
        return guard
    return send_from_directory(PROJECT_ROOT / "frontend", request.path.lstrip("/"))


@app.get("/manifest.webmanifest")
def original_manifest():
    # Browsers fetch manifests without the app's login cookie on some platforms.
    # This file contains only public app metadata; the start_url remains protected.
    return send_from_directory(PROJECT_ROOT / "frontend", "manifest.webmanifest",
                               mimetype="application/manifest+json")


@app.get("/yuling-assets/<path:filename>")
def original_yuling_asset(filename):
    guard = _require_login()
    if guard:
        return guard
    return send_from_directory(PROJECT_ROOT / "frontend" / "vendor", filename)


@app.get("/yuling-media/<path:filename>")
def original_yuling_media(filename):
    guard = _require_login()
    if guard:
        return guard
    if filename not in {
        "chen-dunni-loop.m4v", "chen-dunni-poster.jpg", "ringback.wav",
        "dunni-emotion-happy-v1.mp4", "dunni-emotion-curious-v1.mp4",
        "dunni-emotion-unhappy-v1.mp4", "dunni-emotion-sad-v1.mp4",
        "dunni-emotion-angry-v1.mp4", "dunni-emotion-delighted-v1.mp4",
        "dunni-emotion-playful_face-v1.mp4",
    }:
        abort(404)
    mimetype = ("video/mp4" if filename.endswith((".m4v", ".mp4")) else
                "audio/wav" if filename.endswith(".wav") else "image/jpeg")
    return send_from_directory(PROJECT_ROOT / "frontend" / "media", filename,
                               mimetype=mimetype, conditional=True)


@app.post("/api/login")
def login():
    if not _configured():
        return jsonify(error="登录服务未设置"), 503
    if not _same_origin() or request.mimetype != "application/json" or request.content_length is None or request.content_length > 4096:
        return jsonify(error="请求格式不正确"), 403
    key = request.remote_addr or "unknown"
    now = time.time()
    with _failure_lock:
        recent = _failures[key]
        while recent and now - recent[0] > 900:
            recent.popleft()
        if len(recent) >= 8:
            return jsonify(error="请稍后再试"), 429
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify(error="请求格式不正确"), 400
    username, password = data.get("username"), data.get("password")
    if not isinstance(username, str) or not isinstance(password, str):
        return jsonify(error="帐号或密码不正确"), 401
    canonical_username = username.strip()
    if USERNAME and canonical_username.casefold() == USERNAME.casefold():
        canonical_username = USERNAME
        correct = hmac.compare_digest(password, PASSWORD)
    else:
        matches = [name for name in EXTRA_ACCOUNTS
                   if name.casefold() == canonical_username.casefold()]
        canonical_username = matches[0] if len(matches) == 1 else ""
        correct = bool(canonical_username) and check_password_hash(
            EXTRA_ACCOUNTS[canonical_username], password)
    if not correct:
        with _failure_lock:
            _failures[key].append(now)
        return jsonify(error="帐号或密码不正确"), 401
    with _failure_lock:
        _failures.pop(key, None)
    remember = data.get("remember") is True
    session.clear()
    session.permanent = remember
    session["account"] = canonical_username
    session["credential_tag"] = _credential_tag(canonical_username)
    session["until"] = now + (REMEMBER_SECONDS if remember else SHORT_SECONDS)
    return jsonify(ok=True, remember=remember)


@app.post("/api/logout")
def logout():
    guard = _require_login(api=True)
    if guard:
        return guard
    if not _same_origin():
        return jsonify(error="来源不符"), 403
    session.clear()
    return jsonify(ok=True)


@app.get("/api/v2/account")
def current_account():
    guard = _require_login(api=True)
    if guard:
        return guard
    store = _chat_store()
    profile = _profile_store(store) if store is not None else None
    display_name = (profile.identity() or {}).get("display_name") if profile else None
    return jsonify(account=session["account"], display_name=display_name)


def _bridge(path, body=None, content_type=None):
    headers = {"Host": _bridge_url.netloc, "Origin": BRIDGE}
    if content_type:
        headers["Content-Type"] = content_type
    call = Request(BRIDGE + path, data=body, headers=headers,
                   method="POST" if body is not None else "GET")
    try:
        with urlopen(call, timeout=90) as response:
            return response.status, json.load(response)
    except HTTPError as exc:
        try:
            return exc.code, json.load(exc)
        except (ValueError, UnicodeDecodeError):
            return exc.code, {"error": "本机桥接服务失败"}
    except (URLError, TimeoutError):
        return 503, {"error": "本机桥接服务无法连接"}


def _chat_store():
    try:
        owner_id = session.get("account")
        if not isinstance(owner_id, str) or not owner_id:
            return None
        return ChatStore(os.environ.get("POC_CHAT_DB_PATH", ""), owner_id,
                         legacy_owner_id=os.environ.get("POC_LEGACY_OWNER", USERNAME))
    except (StoreUnavailable, OSError):
        return None


_USAGE_COMPONENTS = {"llm", "stt", "tts", "turn_detector", "interruption"}
_USAGE_META_FIELDS = {"source_status", "cost_status", "input_cached_tokens",
                      "input_cache_creation_tokens", "input_audio_tokens",
                      "output_audio_tokens", "total_requests", "session_duration"}


def _usage_store(owner_id):
    """A Railway receipt must live on its mounted volume, or fail closed."""
    path = os.environ.get("POC_CHAT_DB_PATH", "").strip()
    if not path or not Path(path).is_absolute():
        return None
    if os.environ.get("RAILWAY_SERVICE_ID"):
        mount = os.environ.get("RAILWAY_VOLUME_MOUNT_PATH", "").strip()
        if not mount or not Path(path).resolve().is_relative_to(Path(mount).resolve()):
            return None
    try:
        return ChatStore(path, owner_id,
                         legacy_owner_id=os.environ.get("POC_LEGACY_OWNER", USERNAME))
    except (StoreUnavailable, OSError):
        return None


def _usage_receipt(raw):
    """Validate a signed, strictly metering-only SDK receipt."""
    try:
        item = json.loads(raw)
    except (ValueError, UnicodeDecodeError):
        return None
    required = {"event_id", "owner_id", "call_id", "created_at", "provider",
                "component", "model", "input_tokens", "output_tokens",
                "audio_input_seconds", "audio_output_chars", "estimated_usd_micros",
                "metadata"}
    if not isinstance(item, dict) or set(item) != required:
        return None
    if (not isinstance(item["event_id"], str) or
            not re.fullmatch(r"[0-9a-f]{64}", item["event_id"]) or
            not isinstance(item["owner_id"], str) or
            item["owner_id"] not in ({USERNAME} | set(EXTRA_ACCOUNTS)) or
            not isinstance(item["call_id"], str) or
            not re.fullmatch(r"ailivex-yuling-[A-Za-z0-9_-]{1,115}", item["call_id"]) or
            not isinstance(item["component"], str) or
            item["component"] not in _USAGE_COMPONENTS):
        return None
    for key, limit in (("provider", 100), ("model", 120)):
        value = item[key]
        if not isinstance(value, str) or not value or len(value) > limit or any(
                ord(char) < 32 for char in value):
            return None
    timestamp = item["created_at"]
    if (isinstance(timestamp, bool) or not isinstance(timestamp, (int, float)) or
            not math.isfinite(timestamp) or abs(time.time() - timestamp) > 300):
        return None
    for key in ("input_tokens", "output_tokens", "audio_output_chars"):
        value = item[key]
        if isinstance(value, bool) or not isinstance(value, int) or not 0 <= value <= 10**12:
            return None
    duration = item["audio_input_seconds"]
    if (isinstance(duration, bool) or not isinstance(duration, (int, float)) or
            not math.isfinite(duration) or not 0 <= duration <= 86_400):
        return None
    # Pricing is not configured: no guessed USD cost or credit is accepted.
    if item["estimated_usd_micros"] is not None:
        return None
    metadata = item["metadata"]
    if (not isinstance(metadata, dict) or not set(metadata) <= _USAGE_META_FIELDS or
            metadata.get("source_status") != "livekit_sdk_reported" or
            metadata.get("cost_status") != "unpriced"):
        return None
    for key, value in metadata.items():
        if key in {"source_status", "cost_status"}:
            continue
        if key == "session_duration":
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or not 0 <= value <= 86_400:
                return None
        elif isinstance(value, bool) or not isinstance(value, int) or not 0 <= value <= 10**12:
            return None
    return item


@app.post("/api/yuling/usage/ingest")
def ingest_voice_usage():
    key = os.environ.get("POC_USAGE_HMAC_KEY", "")
    if len(key.encode("utf-8")) < 32:
        return jsonify(error="用量接收未配置"), 503
    if (request.mimetype != "application/json" or request.content_length is None or
            not 0 < request.content_length <= 4096):
        return jsonify(error="请求格式不正确"), 400
    raw = request.get_data(cache=False)
    signature = request.headers.get("X-Usage-Signature", "")
    expected = "sha256=" + hmac.new(key.encode("utf-8"), raw, sha256).hexdigest()
    if not hmac.compare_digest(signature, expected):
        return jsonify(error="签名不正确"), 403
    item = _usage_receipt(raw)
    if item is None:
        return jsonify(error="用量格式不正确"), 400
    store = _usage_store(item["owner_id"])
    if store is None:
        return jsonify(error="用量存储未就绪"), 503
    try:
        changed = store.record_voice_usage(item)
    except StoreUnavailable:
        return jsonify(error="用量事件冲突"), 409
    return jsonify(ok=True, changed=changed), 201 if changed else 200


@app.get("/api/yuling/usage")
def my_voice_usage():
    # The before_request guard already validated the login cookie. A query
    # parameter can never select another owner's usage.
    store = _usage_store(request.environ["poc.authenticated_owner"])
    if store is None:
        return jsonify(error="用量存储未就绪"), 503
    return jsonify(store.voice_usage_summary())


def _companion_root():
    """Optional directory for separately stored profile and skill candidates."""
    value = os.environ.get("POC_COMPANION_ROOT", "").strip()
    root = Path(value).expanduser() if value else None
    if root is None or not root.is_absolute():
        return None
    # Railway's container filesystem is disposable. Do not invite a person to
    # opt into long-term profile capture unless the intended directory is on
    # the attached volume. Local test installations need only an absolute path.
    if os.environ.get("RAILWAY_SERVICE_ID"):
        mount = os.environ.get("RAILWAY_VOLUME_MOUNT_PATH", "").strip()
        if not mount or not root.resolve().is_relative_to(Path(mount).resolve()):
            return None
    return root


def _profile_store(chat_store):
    root = _companion_root()
    if root is None or chat_store is None:
        return None
    if os.environ.get("RAILWAY_SERVICE_ID"):
        mount = Path(os.environ["RAILWAY_VOLUME_MOUNT_PATH"]).resolve()
        if not chat_store.path.resolve().is_relative_to(mount):
            return None
    try:
        return NicoProfileStore(str(chat_store.path), chat_store.owner_id, str(root))
    except (ValueError, OSError):
        return None


def _skill_store():
    root = _companion_root()
    if root is None:
        return None
    try:
        return SkillGrowthStore(str(root / "skills.sqlite3"))
    except (ValueError, OSError):
        return None


def _uuid(value):
    try:
        return str(UUID(value)) if isinstance(value, str) else None
    except ValueError:
        return None


def _v2_guard(write=False):
    guard = _require_login(api=True)
    if guard:
        return guard
    if write and not _same_origin():
        return jsonify(error="来源不符"), 403
    return None


@app.get("/api/v2/tools/time")
def time_tool_info():
    guard = _v2_guard()
    if guard:
        return guard
    return jsonify(tool="time_convert_local", zones=ZONES,
                   permissions=[], cost="free", status="local_ready")


@app.get("/api/v2/tools/catalog")
def tool_catalog():
    """Only enabled local capabilities get an action; candidates stay inert."""
    guard = _v2_guard()
    if guard:
        return guard
    page_read_ready = (os.environ.get("POC_PUBLIC_WEB_READ_ENABLED") == "1" and
                       _chat_store() is not None)
    return jsonify(groups=[
        {"title": "已可使用", "items": [
            {"name": "时区换算", "state": "ready", "detail": "本机换算，不需要外部帐号",
             "href": "/tools"},
            {"name": "公开网页读取", "state": "ready" if page_read_ready else "waiting",
             "detail": ("在对话中发送公开 HTTPS 网址，再点读取此网页" if page_read_ready
                        else "后端读取功能尚未启用"),
             **({"href": "/chat"} if page_read_ready else {})},
        ]},
        {"title": "生活服务 · 等待整合", "items": [
            {"name": "和风天气", "state": "waiting", "detail": "待专属 Key、来源标示与费用硬上限"},
            {"name": "高德地图", "state": "waiting", "detail": "待 Web Key 与服务授权确认"},
            {"name": "支付宝", "state": "waiting", "detail": "待商户资格与个人资料授权确认"},
            {"name": "美团", "state": "waiting", "detail": "待合适的个人服务接口与授权"},
            {"name": "QQ 邮箱", "state": "waiting",
             "detail": "待用户开启 IMAP、提供专用授权码及只读接入验收"},
            {"name": "12306", "state": "unavailable", "detail": "暂不提供抓票、代购或帐号连接"},
        ]},
        {"title": "工作服务 · 等待整合", "items": [
            {"name": "飞书", "state": "waiting", "detail": "待本人授权与唯读范围确认"},
            {"name": "钉钉", "state": "waiting", "detail": "待组织管理员授权与唯读范围确认"},
            {"name": "WPS 365", "state": "waiting", "detail": "待企业应用及本人授权"},
            {"name": "企业微信", "state": "waiting", "detail": "待企业授权与接口范围确认"},
            {"name": "腾讯会议", "state": "waiting", "detail": "待本人授权与唯读范围确认"},
        ]},
        {"title": "趣味探索 · 等待整合", "items": [
            {"name": "B站投稿回顾", "state": "waiting",
             "detail": "待开放平台应用、本人授权与只读稿件权限"},
            {"name": "抖音作品互动", "state": "waiting",
             "detail": "待应用审核、本人授权与作品数据权限"},
            {"name": "快手漫画脸", "state": "waiting",
             "detail": "待图像授权、接口资格与费用确认"},
            {"name": "Lichess 每日棋题", "state": "waiting",
             "detail": "待中国大陆连线、限流与展示验收"},
            {"name": "NASA 每日宇宙图片", "state": "waiting",
             "detail": "待新版接口、图片许可与中国大陆连线核对"},
        ]},
    ])


@app.post("/api/v2/tools/time")
def time_tool_convert():
    guard = _v2_guard(write=True)
    if guard:
        return guard
    if (request.mimetype != "application/json" or not request.content_length or
            request.content_length > 256):
        return jsonify(error="时间请求格式不正确"), 415
    data = request.get_json(silent=True)
    if not isinstance(data, dict) or set(data) != {"time", "source_timezone", "target_timezone"}:
        return jsonify(error="时间请求字段不正确"), 400
    try:
        result = convert_time(data["time"], data["source_timezone"], data["target_timezone"])
    except TimeToolError as exc:
        return jsonify(error=str(exc)), 400
    return jsonify(result=result)


@app.get("/api/v2/profile")
def companion_profile():
    guard = _v2_guard()
    if guard:
        return guard
    profile = _profile_store(_chat_store())
    if profile is None:
        return jsonify(error="个人资料存储尚未配置"), 503
    return jsonify(identity=profile.identity(),
                   consent_enabled=profile.consent_status(),
                   consent_decided=profile.consent_decided(),
                   facts=profile.facts(), account_seed=profile.seed_facts())


@app.patch("/api/v2/profile/account-seed/<field>")
def correct_account_seed(field):
    guard = _v2_guard(write=True)
    if guard:
        return guard
    if request.mimetype != "application/json" or not request.content_length or request.content_length > 512:
        return jsonify(error="更正格式不正确"), 415
    data = request.get_json(silent=True)
    if not isinstance(data, dict) or set(data) != {"value"}:
        return jsonify(error="更正格式不正确"), 400
    profile = _profile_store(_chat_store())
    if profile is None:
        return jsonify(error="个人资料存储尚未配置"), 503
    try:
        item = profile.correct_seed_fact(field, data["value"])
    except ValueError as exc:
        return jsonify(error=str(exc)), 400
    return jsonify(item=item)


@app.post("/api/v2/profile/consent")
def set_companion_profile_consent():
    guard = _v2_guard(write=True)
    if guard:
        return guard
    if (request.mimetype != "application/json" or not request.content_length or
            request.content_length > 256):
        return jsonify(error="同意设置格式不正确"), 415
    data = request.get_json(silent=True)
    if not isinstance(data, dict) or set(data) != {"enabled"} or type(data["enabled"]) is not bool:
        return jsonify(error="同意设置格式不正确"), 400
    profile = _profile_store(_chat_store())
    if profile is None:
        return jsonify(error="个人资料存储尚未配置"), 503
    enabled = profile.set_consent(data["enabled"])
    return jsonify(consent_enabled=enabled, consent_decided=True)


@app.delete("/api/v2/profile/facts/<fact_id>")
def delete_companion_profile_fact(fact_id):
    guard = _v2_guard(write=True)
    if guard:
        return guard
    fact_id = _uuid(fact_id)
    if fact_id is None:
        return jsonify(error="记忆编号不正确"), 400
    profile = _profile_store(_chat_store())
    if profile is None:
        return jsonify(error="个人资料存储尚未配置"), 503
    if not profile.delete_fact(fact_id):
        return jsonify(error="记忆不存在"), 404
    return jsonify(ok=True)


@app.get("/api/v2/skill-requests")
def companion_skill_requests():
    guard = _v2_guard()
    if guard:
        return guard
    skills = _skill_store()
    if skills is None:
        return jsonify(error="技能候选存储尚未配置"), 503
    return jsonify(candidates=skills.list(session["account"]))


@app.get("/api/v2/version")
def deployed_version():
    guard = _v2_guard()
    if guard:
        return guard
    # Railway provides these values for the running deployment.  A Docker
    # build does not include .git, so never display the checkout's local SHA.
    deployment = os.environ.get("RAILWAY_DEPLOYMENT_ID", "").strip()
    commit = os.environ.get("RAILWAY_GIT_COMMIT_SHA", "").strip().lower()
    deployment = deployment if _uuid(deployment) else ""
    commit = commit if re.fullmatch(r"[0-9a-f]{40}", commit) else ""
    # Show the code revision people can quote back to us. Keep the deployment
    # ID as the refresh cache key, since a redeploy of the same commit is new.
    version = commit[:7] if commit else ("v" + deployment[:8] if deployment else None)
    return jsonify(version=version,
                   cache_key=deployment or commit or None,
                   commit=commit[:8] if commit else None)


@app.get("/api/v2/status")
def conversation_status():
    guard = _v2_guard()
    if guard:
        return guard
    store = _chat_store()
    store_path = os.path.realpath(str(store.path)) if store else ""
    temporary = store_path.startswith(("/tmp/", "/private/tmp/"))
    # A configured absolute path does not prove that Railway mounted durable
    # storage there. Report persistence only after a separate restore check.
    railway_without_volume = bool(os.environ.get("RAILWAY_SERVICE_ID")) and not bool(
        os.environ.get("RAILWAY_VOLUME_MOUNT_PATH"))
    persistence = ("unavailable" if store is None else
                   "ephemeral" if temporary or railway_without_volume else
                   "path_configured_unverified")
    return jsonify(text_model=bool(os.environ.get("MINIMAX_API_KEY", "").strip()),
                   persistent_history=False,
                   memory_persistence=persistence,
                   image_upload=store is not None,
                   image_understanding=bool(os.environ.get("MINIMAX_API_KEY", "").strip()),
                   image_understanding_verified=False, vector_search=False,
                   sourced_knowledge=store is not None,
                   graph_reasoning=False, hermes=False)


@app.post("/api/v2/upload-image")
def upload_chat_image():
    guard = _v2_guard(write=True)
    if guard:
        return guard
    store = _chat_store()
    if store is None:
        return jsonify(error="持久化存储尚未配置"), 503
    file = request.files.get("file")
    if file is None:
        return jsonify(error="请选择图片"), 400
    data = file.stream.read(4_000_001)
    if not data or len(data) > 4_000_000:
        return jsonify(error="图片大小需在 4 MB 内"), 413
    try:
        from PIL import Image, ImageOps, UnidentifiedImageError
        with Image.open(io.BytesIO(data)) as image:
            if image.width * image.height > 20_000_000:
                return jsonify(error="图片像素过大"), 413
            image.verify()
        with Image.open(io.BytesIO(data)) as image:
            if image.width * image.height > 20_000_000:
                return jsonify(error="图片像素过大"), 413
            mime = {"PNG": "image/png", "JPEG": "image/jpeg", "WEBP": "image/webp"}.get(image.format)
            if mime is None:
                return jsonify(error="仅支持 PNG、JPEG、WebP"), 415
            # Re-encode pixels into a fresh image. This removes EXIF/GPS and
            # other embedded metadata before any bytes enter the database.
            image = ImageOps.exif_transpose(image)
            image.thumbnail((2048, 2048), Image.Resampling.LANCZOS)
            mode = "RGBA" if image.mode in {"RGBA", "LA"} else "RGB"
            fresh = Image.new(mode, image.size)
            fresh.paste(image.convert(mode))
            cleaned = io.BytesIO()
            if mime == "image/jpeg":
                fresh.convert("RGB").save(cleaned, "JPEG", quality=85)
            elif mime == "image/webp":
                fresh.save(cleaned, "WEBP", quality=85)
            else:
                fresh.save(cleaned, "PNG", optimize=True)
            data = cleaned.getvalue()
    except (UnidentifiedImageError, OSError, ValueError, Image.DecompressionBombError):
        return jsonify(error="图片格式不正确"), 400
    if len(data) > 4_000_000:
        return jsonify(error="处理后图片超过 4 MB"), 413
    iid = store.put_image(data, mime)
    return jsonify(image_id=iid, url=f"/api/v2/image/{iid}", mime_type=mime,
                   image_understanding="待发送并验证"), 201


@app.get("/api/v2/image/<image_id>")
def chat_image(image_id):
    guard = _v2_guard()
    if guard:
        return guard
    store = _chat_store()
    if store is None:
        return jsonify(error="持久化存储尚未配置"), 503
    iid = _uuid(image_id)
    row = store.get_image(iid) if iid else None
    if row is None:
        abort(404)
    return send_file(io.BytesIO(row[1]), mimetype=row[0], conditional=False,
                     download_name="image")


DOCUMENT_MIMES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".txt": "text/plain", ".md": "text/markdown", ".csv": "text/csv",
    ".doc": "application/octet-stream", ".xls": "application/octet-stream",
    ".ppt": "application/octet-stream", ".wps": "application/octet-stream",
    ".et": "application/octet-stream", ".dps": "application/octet-stream",
    ".ofd": "application/octet-stream",
}
UNPARSED_DOCUMENT_EXTENSIONS = {".doc", ".xls", ".ppt", ".wps", ".et", ".dps", ".ofd"}
OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _valid_unparsed_document(ext, data):
    if ext in {".doc", ".xls", ".ppt"}:
        return data.startswith(OLE_SIGNATURE)
    if ext in {".wps", ".et", ".dps"}:
        return data.startswith(OLE_SIGNATURE) or data.startswith(b"PK\x03\x04")
    if ext == ".ofd" and data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                names = archive.namelist()
                return (len(names) <= 1000 and
                        sum(item.file_size for item in archive.infolist()) <= 20_000_000 and
                        "OFD.xml" in names)
        except zipfile.BadZipFile:
            return False
    return False


@app.post("/api/v2/upload-document")
def upload_chat_document():
    guard = _v2_guard(write=True)
    if guard:
        return guard
    store = _chat_store()
    if store is None:
        return jsonify(error="文件存储尚未配置"), 503
    file = request.files.get("file")
    if file is None:
        return jsonify(error="请选择文件"), 400
    filename = (file.filename or "").replace("\\", "/").rsplit("/", 1)[-1].strip()
    if not filename or len(filename) > 150 or any(ord(ch) < 32 for ch in filename):
        return jsonify(error="文件名不正确"), 400
    ext = Path(filename).suffix.lower()
    if ext not in DOCUMENT_MIMES:
        return jsonify(error="文件格式暂不支持上传"), 415
    data = file.stream.read(MAX_FILE + 1)
    if not data or len(data) > MAX_FILE:
        return jsonify(error="文件大小需在 5 MB 内"), 413
    if ext in UNPARSED_DOCUMENT_EXTENSIONS:
        if not _valid_unparsed_document(ext, data):
            return jsonify(error="文件内容与格式不符或暂无法验证"), 400
        extracted = {"kind": ext[1:].upper(), "summary": "", "characters": 0}
        status = "unsupported_extract"
    else:
        try:
            extracted = extract_document(filename, data)
        except (ValueError, zipfile.BadZipFile, ET.ParseError, ImportError, UnicodeDecodeError) as exc:
            return jsonify(error=str(exc) or "无法读取文件"), 400
        except Exception:
            return jsonify(error="无法抽取此文件的文字"), 422
        status = "text_preview" if extracted["characters"] else "no_text"
    document = store.put_document(filename, extracted["kind"], DOCUMENT_MIMES[ext],
                                  data, extracted["summary"] if status == "text_preview" else "",
                                  status)
    return jsonify(document_id=document["id"], document=document,
                   characters=extracted["characters"],
                   note=("已保存，可下载；此格式暂不支持内容解析。" if status == "unsupported_extract"
                         else "仅抽取前 2500 字供本轮对话；图片、扫描页与版式未被理解。")), 201


@app.get("/api/v2/document/<document_id>")
def chat_document(document_id):
    guard = _v2_guard()
    if guard:
        return guard
    store = _chat_store()
    if store is None:
        return jsonify(error="文件存储尚未配置"), 503
    did = _uuid(document_id)
    row = store.get_document(did, include_data=True) if did else None
    if row is None:
        abort(404)
    return send_file(io.BytesIO(row["data"]), mimetype=row["mime_type"],
                     as_attachment=True, download_name=row["filename"], conditional=False)


@app.get("/api/v2/history")
def conversation_history():
    guard = _v2_guard()
    if guard:
        return guard
    store = _chat_store()
    if store is None:
        return jsonify(error="持久化存储尚未配置"), 503
    sid = _uuid(request.args.get("session_id"))
    if request.args.get("session_id") and not sid:
        return jsonify(error="会话编号不正确"), 400
    if sid:
        if not store.has_session(sid):
            return jsonify(error="会话不存在"), 404
        return jsonify(session_id=sid, events=store.history(sid))
    return jsonify(sessions=store.recent_sessions(), events=[])


@app.post("/api/v2/page-read")
def read_conversation_page():
    """Opt-in read of an existing user link card; never feeds text to the model."""
    guard = _v2_guard(write=True)
    if guard:
        return guard
    if os.environ.get("POC_PUBLIC_WEB_READ_ENABLED") != "1":
        return jsonify(error="网页读取尚未启用"), 503
    if (request.mimetype != "application/json" or not request.content_length or
            request.content_length > 1024):
        return jsonify(error="网页读取请求格式不正确"), 415
    data = request.get_json(silent=True)
    if not isinstance(data, dict) or set(data) != {"source_event_id"}:
        return jsonify(error="只需提供来源消息编号"), 400
    source_event_id = _uuid(data["source_event_id"])
    if not source_event_id:
        return jsonify(error="来源消息编号不正确"), 400
    store = _chat_store()
    if store is None:
        return jsonify(error="对话存储尚未配置"), 503
    source_url = store.page_read_source(source_event_id)
    if source_url is None:
        return jsonify(error="来源消息不存在或没有网址"), 404
    try:
        _, canonical_url = _validated_url(source_url)
    except PageReadError:
        return jsonify(error="仅支持公开 HTTPS 网页"), 400
    if not _page_read_lock.acquire(blocking=False):
        return jsonify(error="网页读取忙，请稍后再试"), 429
    try:
        try:
            page_id, remaining = store.reserve_page_read(source_event_id, canonical_url)
        except StoreUnavailable:
            return jsonify(error="来源消息不存在或没有网址"), 404
        except PageReadQuotaExceeded:
            return jsonify(error="今日网页读取次数已用完"), 429
        try:
            result = read_public_page(source_url)
        except PageReadError:
            store.finish_page_read(page_id)
            return jsonify(page_read=store.page_read(page_id),
                           remaining_today=remaining,
                           error="网页暂时无法读取"), 502
        except Exception:
            # Unexpected parser/transport failures still consume the attempt
            # and leave an honest failed artifact, without exposing details.
            store.finish_page_read(page_id)
            return jsonify(page_read=store.page_read(page_id),
                           remaining_today=remaining,
                           error="网页暂时无法读取"), 502
        store.finish_page_read(page_id, result)
        return jsonify(page_read=store.page_read(page_id),
                       remaining_today=remaining), 201
    finally:
        _page_read_lock.release()


@app.get("/api/v2/page-read/<page_id>")
def get_conversation_page_read(page_id):
    guard = _v2_guard()
    if guard:
        return guard
    if os.environ.get("POC_PUBLIC_WEB_READ_ENABLED") != "1":
        return jsonify(error="网页读取尚未启用"), 503
    store = _chat_store()
    if store is None:
        return jsonify(error="对话存储尚未配置"), 503
    artifact = store.page_read(_uuid(page_id)) if _uuid(page_id) else None
    if artifact is None:
        return jsonify(error="网页读取记录不存在"), 404
    return jsonify(page_read=artifact)


@app.post("/api/v2/call-transcript")
def save_call_transcript():
    """Import final LiveKit segments observed by this authenticated call browser."""
    guard = _v2_guard(write=True)
    if guard:
        return guard
    if request.mimetype != "application/json" or not request.content_length or request.content_length > MAX_JSON:
        return jsonify(error="逐字稿格式不正确"), 415
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify(error="逐字稿格式不正确"), 400
    call_id = data.get("call_id")
    speaker_identity = data.get("speaker_identity")
    segment_id = data.get("segment_id")
    role = data.get("role")
    text_value = data.get("text")
    if (not isinstance(call_id, str) or not 1 <= len(call_id) <= 128 or
            not call_id.startswith("ailivex-yuling-") or
            not re.fullmatch(r"[A-Za-z0-9_-]+", call_id)):
        return jsonify(error="通话编号不正确"), 400
    if (not isinstance(speaker_identity, str) or not 1 <= len(speaker_identity) <= 128 or
            not re.fullmatch(r"[A-Za-z0-9_.:-]+", speaker_identity)):
        return jsonify(error="说话者编号不正确"), 400
    if (not isinstance(segment_id, str) or not 1 <= len(segment_id) <= 128 or
            not re.fullmatch(r"[A-Za-z0-9_.:-]+", segment_id)):
        return jsonify(error="片段编号不正确"), 400
    if role not in {"user", "assistant"}:
        return jsonify(error="说话者类型不正确"), 400
    if not isinstance(text_value, str) or not 1 <= len(text_value.strip()) <= 4000:
        return jsonify(error="逐字稿长度不正确"), 400
    sid_raw = data.get("session_id")
    sid = _uuid(sid_raw) if sid_raw else None
    if sid_raw and not sid:
        return jsonify(error="会话编号不正确"), 400
    store = _chat_store()
    if store is None:
        return jsonify(error="对话存储尚未配置"), 503
    try:
        event, created = store.add_call_transcript(
            sid, call_id, speaker_identity, segment_id, role, text_value.strip())
    except StoreUnavailable:
        return jsonify(error="会话不存在"), 404
    return jsonify(session_id=event["session_id"], event=event,
                   created=created), 201 if created else 200


@app.get("/api/v2/calls")
def saved_calls():
    guard = _v2_guard()
    if guard:
        return guard
    store = _chat_store()
    if store is None:
        return jsonify(error="对话存储尚未配置"), 503
    return jsonify(calls=store.recent_calls(), audio_saved=False,
                   source="final_browser_transcript")


@app.get("/api/v2/calls/<call_id>")
def saved_call(call_id):
    guard = _v2_guard()
    if guard:
        return guard
    if not re.fullmatch(r"ailivex-yuling-[A-Za-z0-9_-]{1,113}", call_id):
        abort(404)
    store = _chat_store()
    if store is None:
        return jsonify(error="对话存储尚未配置"), 503
    detail = store.call_transcript(call_id)
    if detail is None:
        abort(404)
    return jsonify(detail)


@app.get("/api/v2/call-reflections")
def saved_call_reflections():
    guard = _v2_guard()
    if guard:
        return guard
    store = _chat_store()
    if store is None:
        return jsonify(error="对话存储尚未配置"), 503
    return jsonify(items=store.recent_call_reflections(),
                   source="saved_final_call_transcript")


@app.post("/api/v2/calls/<call_id>/reflection")
def reflect_on_call(call_id):
    guard = _v2_guard(write=True)
    if guard:
        return guard
    if not re.fullmatch(r"ailivex-yuling-[A-Za-z0-9_-]{1,113}", call_id):
        abort(404)
    if request.mimetype != "application/json" or not request.content_length or request.content_length > 256:
        return jsonify(error="请求格式不正确"), 415
    data = request.get_json(silent=True)
    if (not isinstance(data, dict) or set(data) - {"retry"} or
            type(data.get("retry", False)) is not bool):
        return jsonify(error="请求格式不正确"), 400
    store = _chat_store()
    if store is None:
        return jsonify(error="对话存储尚未配置"), 503
    claim, events = store.claim_call_reflection(
        call_id, PROMPT_VERSION, retry=data.get("retry") is True)
    if claim is None:
        abort(404)
    if claim == "existing":
        return jsonify(item=store.call_reflection(call_id), created=False)
    if not substantial_call(events):
        store.finish_call_reflection(call_id, "skipped")
        return jsonify(item=store.call_reflection(call_id), created=True), 201
    def observe_reflection_usage(observation):
        try:
            store.record_text_model_usage(scope="call_reflection", scope_id=call_id,
                                          observation=observation)
        except (StoreUnavailable, OSError):
            app.logger.exception("call reflection usage recording failed")
    try:
        note, model = generate_reflection(events, usage_observer=observe_reflection_usage)
    except (ProviderError, ValueError, TypeError):
        store.finish_call_reflection(call_id, "failed")
        return jsonify(item=store.call_reflection(call_id),
                       error="通话后心声未生成，请稍后手动重试"), 502
    store.finish_call_reflection(call_id, "generated" if note else "skipped",
                                 note, model)
    return jsonify(item=store.call_reflection(call_id), created=True), 201


@app.post("/api/v2/chat")
def conversation_chat():
    guard = _v2_guard(write=True)
    if guard:
        return guard
    if request.mimetype != "application/json" or not request.content_length or request.content_length > MAX_JSON:
        return jsonify(error="文字格式不正确"), 415
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify(error="文字格式不正确"), 400
    text_value = data.get("text", "")
    if not isinstance(text_value, str) or len(text_value) > 4000:
        return jsonify(error="文字最多 4000 字"), 400
    text_value = text_value.strip()
    image_id = _uuid(data.get("image_id")) if data.get("image_id") else None
    if data.get("image_id") and not image_id:
        return jsonify(error="图片编号不正确"), 400
    document_id = _uuid(data.get("document_id")) if data.get("document_id") else None
    if data.get("document_id") and not document_id:
        return jsonify(error="文件编号不正确"), 400
    page_read_id = _uuid(data.get("page_read_id")) if data.get("page_read_id") else None
    if data.get("page_read_id") and not page_read_id:
        return jsonify(error="网页读取编号不正确"), 400
    if page_read_id and os.environ.get("POC_PUBLIC_WEB_READ_ENABLED") != "1":
        return jsonify(error="网页读取尚未启用"), 503
    link_url = data.get("url")
    if link_url is not None:
        if (not isinstance(link_url, str) or len(link_url) > 2048 or
                urlparse(link_url).scheme not in {"http", "https"} or not urlparse(link_url).hostname):
            return jsonify(error="网址格式不正确"), 400
    if not text_value and not image_id and not document_id and not link_url and not page_read_id:
        return jsonify(error="请输入文字、图片、文件或网址"), 400
    store = _chat_store()
    if store is None:
        return jsonify(error="持久化存储尚未配置"), 503
    if not os.environ.get("MINIMAX_API_KEY", "").strip():
        return jsonify(error="对话模型尚未配置"), 503
    image = store.get_image(image_id) if image_id else None
    if image_id and image is None:
        return jsonify(error="图片不存在"), 404
    document = store.get_document(document_id) if document_id else None
    if document_id and document is None:
        return jsonify(error="文件不存在"), 404
    sid_raw = data.get("session_id")
    sid = _uuid(sid_raw) if sid_raw else None
    if sid_raw and not sid:
        return jsonify(error="会话编号不正确"), 400
    if page_read_id and sid is None:
        return jsonify(error="引用网页需要原会话编号"), 400
    if sid and not store.has_session(sid):
        return jsonify(error="会话不存在"), 404
    page_context = store.usable_page_context(page_read_id, sid) if page_read_id else None
    if page_read_id and page_context is None:
        return jsonify(error="网页读取记录不存在或已过期"), 404
    sid = sid or store.create_session()
    history = store.history(sid, limit=20)
    try:
        user_event = store.add_event(sid, "user", text_value, image_id=image_id,
                                     link_url=link_url, source_kind="user_explicit",
                                     page_read_id=page_read_id, document_id=document_id)
    except StoreUnavailable:
        return jsonify(error="网页读取记录不存在或已过期"), 404
    # Profile candidates come only from the authenticated user's own saved
    # event, and only after that account has explicitly opted in. They are
    # candidates, not established personal facts or model instructions.
    profile = _profile_store(store)
    if profile is not None and profile.consent_status() and text_value:
        try:
            profile.propose_from_user_event(user_event["id"])
        except Exception:
            app.logger.exception("profile candidate capture failed")
    need = explicit_skill_need(text_value)
    if need:
        skills = _skill_store()
        if skills is not None:
            try:
                skills.propose(owner_id=store.owner_id,
                               source_event_id=user_event["id"], user_need=need,
                               name=need[:120], kind="local_feature",
                               source_url=None, license_id="unknown",
                               requested_scopes=[], cost="unknown")
            except Exception:
                app.logger.exception("skill candidate capture failed")
    def observe_chat_usage(observation):
        try:
            store.record_text_model_usage(scope="chat", scope_id=user_event["id"],
                                          observation=observation)
        except (StoreUnavailable, OSError):
            app.logger.exception("text model usage recording failed")
    try:
        answer = reply_to(text_value, history, memories_for_turn(store, text_value), image=image,
                          link_url=link_url, page_context=page_context,
                          document_context=document, owner_id=store.owner_id,
                          display_name=(profile.identity() or {}).get("display_name")
                          if profile is not None else None,
                          profile_candidates=profile.facts(limit=4)
                          if profile is not None and profile.consent_status() else (),
                          account_seed=profile.seed_facts() if profile is not None else {},
                          usage_observer=observe_chat_usage)
    except ProviderError as exc:
        return jsonify(error=str(exc), session_id=sid, user_event=user_event), 502
    assistant_event = store.add_event(sid, "assistant", answer,
                                      source_kind="model_generated",
                                      page_read_id=page_read_id)
    return jsonify(session_id=sid, reply=answer,
                   user_event=user_event, assistant_event=assistant_event)


@app.get("/api/v2/memories")
def explicit_memories():
    guard = _v2_guard()
    if guard:
        return guard
    store = _chat_store()
    if store is None:
        return jsonify(error="持久化存储尚未配置"), 503
    return jsonify(memories=store.memories())


@app.post("/api/v2/memories")
def save_explicit_memory():
    guard = _v2_guard(write=True)
    if guard:
        return guard
    data = request.get_json(silent=True) if request.mimetype == "application/json" else None
    if not isinstance(data, dict) or not isinstance(data.get("text"), str) or not 1 <= len(data["text"].strip()) <= 500:
        return jsonify(error="记忆文字需在 1 至 500 字内"), 400
    store = _chat_store()
    if store is None:
        return jsonify(error="持久化存储尚未配置"), 503
    mid = store.add_memory(data["text"].strip())
    return jsonify(id=mid, source_kind="user_explicit"), 201


@app.delete("/api/v2/memories/<memory_id>")
def forget_explicit_memory(memory_id):
    guard = _v2_guard(write=True)
    if guard:
        return guard
    store = _chat_store()
    if store is None:
        return jsonify(error="持久化存储尚未配置"), 503
    mid = _uuid(memory_id)
    if not mid or not store.delete_memory(mid):
        return jsonify(error="记忆不存在"), 404
    return jsonify(ok=True)


@app.get("/api/v2/knowledge")
def list_knowledge_items():
    guard = _v2_guard()
    if guard:
        return guard
    store = _chat_store()
    if store is None:
        return jsonify(error="知识存储尚未配置"), 503
    return jsonify(items=store.knowledge_items())


@app.post("/api/v2/knowledge")
def save_knowledge_item():
    """Save a user-supplied note linked to its original conversation event."""
    guard = _v2_guard(write=True)
    if guard:
        return guard
    if request.mimetype != "application/json" or not request.content_length or request.content_length > MAX_JSON:
        return jsonify(error="知识格式不正确"), 415
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify(error="知识格式不正确"), 400
    title, content = data.get("title"), data.get("content")
    source_event_id = _uuid(data.get("source_event_id"))
    if not isinstance(title, str) or not 1 <= len(title.strip()) <= 120:
        return jsonify(error="标题需在 1 至 120 字内"), 400
    if not isinstance(content, str) or not 1 <= len(content.strip()) <= 4000:
        return jsonify(error="内容需在 1 至 4000 字内"), 400
    if not source_event_id:
        return jsonify(error="请提供本账号的来源消息"), 400
    store = _chat_store()
    if store is None:
        return jsonify(error="知识存储尚未配置"), 503
    try:
        item = store.add_knowledge_item(title.strip(), content.strip(),
                                        source_event_id, source_kind="user_explicit")
    except StoreUnavailable:
        return jsonify(error="来源消息不存在"), 404
    return jsonify(item=item), 201


@app.delete("/api/v2/knowledge/<item_id>")
def delete_knowledge_item(item_id):
    guard = _v2_guard(write=True)
    if guard:
        return guard
    store = _chat_store()
    if store is None:
        return jsonify(error="知识存储尚未配置"), 503
    kid = _uuid(item_id)
    if not kid or not store.delete_knowledge_item(kid):
        return jsonify(error="知识条目不存在"), 404
    return jsonify(ok=True)


@app.get("/api/v2/relations")
def list_relation_candidates():
    guard = _v2_guard()
    if guard:
        return guard
    store = _chat_store()
    if store is None:
        return jsonify(error="关系存储尚未配置"), 503
    return jsonify(items=store.relation_candidates(), status="candidate_only")


@app.post("/api/v2/relations")
def save_relation_candidate():
    """Record an explicitly sourced possibility, not an established cause."""
    guard = _v2_guard(write=True)
    if guard:
        return guard
    if request.mimetype != "application/json" or not request.content_length or request.content_length > MAX_JSON:
        return jsonify(error="关系格式不正确"), 415
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify(error="关系格式不正确"), 400
    subject_id = _uuid(data.get("subject_id"))
    object_id = _uuid(data.get("object_id"))
    source_event_id = _uuid(data.get("source_event_id"))
    predicate = data.get("predicate")
    if not subject_id or not object_id or not source_event_id:
        return jsonify(error="请提供本账号的知识条目与来源消息"), 400
    if not isinstance(predicate, str) or not 1 <= len(predicate.strip()) <= 120:
        return jsonify(error="关系描述需在 1 至 120 字内"), 400
    store = _chat_store()
    if store is None:
        return jsonify(error="关系存储尚未配置"), 503
    try:
        relation = store.add_relation_candidate(subject_id, predicate.strip(),
                                                object_id, source_event_id)
    except StoreUnavailable:
        return jsonify(error="知识条目或来源消息不存在"), 404
    return jsonify(item=relation), 201


@app.delete("/api/v2/relations/<relation_id>")
def delete_relation_candidate(relation_id):
    guard = _v2_guard(write=True)
    if guard:
        return guard
    store = _chat_store()
    if store is None:
        return jsonify(error="关系存储尚未配置"), 503
    rid = _uuid(relation_id)
    if not rid or not store.delete_relation_candidate(rid):
        return jsonify(error="关系候选不存在"), 404
    return jsonify(ok=True)


def extract_document(filename, data):
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise ValueError("PDF 格式不正确")
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data), strict=False)
        if reader.is_encrypted:
            raise ValueError("加密 PDF 尚不支持")
        if len(reader.pages) > 20:
            raise ValueError("PDF 最多 20 页")
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
        kind = "PDF"
    elif ext in {".docx", ".pptx", ".xlsx"}:
        kind = {".docx": "Word", ".pptx": "PowerPoint", ".xlsx": "Excel"}[ext]
        parts = []
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = archive.namelist()
            if len(names) > 1000 or sum(x.file_size for x in archive.infolist()) > 20_000_000:
                raise ValueError("文件内容过大")
            if "[Content_Types].xml" not in names:
                raise ValueError("Office 文件格式不正确")
            if ext == ".docx":
                targets = [n for n in names if re.fullmatch(r"word/(document|header\d+|footer\d+)\.xml", n)]
            elif ext == ".pptx":
                targets = sorted((n for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                                 key=lambda n: int(re.search(r"slide(\d+)", n).group(1)))
            else:
                targets = []
            if ext == ".xlsx":
                from openpyxl import load_workbook
                workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True,
                                         keep_links=False)
                try:
                    for sheet in workbook.worksheets[:10]:
                        parts.append("[" + sheet.title[:80] + "]")
                        for row in sheet.iter_rows(max_row=min(500, sheet.max_row or 500),
                                                   max_col=min(30, sheet.max_column or 30),
                                                   values_only=True):
                            cells = [str(value)[:200] for value in row if value is not None]
                            if cells:
                                parts.append(" | ".join(cells))
                        if (sheet.max_row or 0) > 500 or (sheet.max_column or 0) > 30:
                            parts.append("…（仅读取前 500 行、30 列）")
                finally:
                    workbook.close()
            for name in targets[:100]:
                raw = archive.read(name)
                if len(raw) > 2_000_000:
                    raise ValueError("单页内容过大")
                root = ET.fromstring(raw)
                parts.append(" ".join(node.text.strip() for node in root.iter()
                                      if node.tag in {"{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t",
                                                      "{http://schemas.openxmlformats.org/drawingml/2006/main}t"}
                                      and node.text and node.text.strip()))
        text = "\n".join(parts)
    elif ext in {".txt", ".md", ".csv"}:
        kind = {".txt": "文本", ".md": "Markdown", ".csv": "CSV"}[ext]
        if b"\x00" in data[:4096]:
            raise ValueError("文本文件格式不正确")
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("gb18030")
    else:
        raise ValueError("支持 PDF、DOCX、PPTX、XLSX、TXT、MD、CSV；DOC、WPS、OFD 等格式请先转为 PDF 或 DOCX")
    clean = re.sub(r"[ \t]+", " ", text).strip()
    if not clean:
        return {"kind": kind, "summary": "未抽出文字；扫描图片或图表内容尚未支持。", "characters": 0}
    preview = clean[:2500]
    if len(clean) > 2500:
        preview += "…（仅显示前 2500 字）"
    return {"kind": kind, "summary": preview, "characters": len(clean),
            "note": "这是文字抽取预览，尚未经模型理解或摘要。"}


@app.get("/api/capabilities")
def capabilities():
    guard = _require_login(api=True)
    if guard:
        return guard
    ready = False
    if BRIDGE:
        status, payload = _bridge("/api/status")
        ready = status == 200 and bool(payload.get("ready"))
    vision_ready = ready and os.path.exists("/Library/Developer/CommandLineTools/SDKs/MacOSX26.5.sdk")
    return jsonify(
        text=ready, turn_voice=ready, live_voice=False,
        image=vision_ready, files=True, sound_events=True,
        image_mode="本机 OCR 与粗分类" if vision_ready else "未接入",
        file_mode="PDF/DOCX/PPTX 文字抽取预览",
        mode="local_bridge" if BRIDGE else "railway_preview",
        reason="本机 Hermes/Whisper 已连接" if ready else "云端尚未接入对话模型；生活声辨识在手机端执行",
    )


@app.post("/api/analyze")
def analyze():
    guard = _require_login(api=True)
    if guard:
        return guard
    if not _same_origin():
        return jsonify(error="来源不符"), 403
    file = request.files.get("file")
    if file is None:
        return jsonify(error="请上传一个文件"), 400
    data = file.stream.read(MAX_FILE + 1)
    if not 0 < len(data) <= MAX_FILE:
        return jsonify(error="文件大小需在 5 MB 内"), 413
    content_type = (file.content_type or "").lower()
    if content_type in {"image/png", "image/jpeg", "image/webp"}:
        if not BRIDGE:
            return jsonify(error="此网址尚未接入图片辨识"), 503
        if len(data) > 4_000_000:
            return jsonify(error="图片需在 4 MB 内"), 413
        status, result = _bridge("/api/analyze", data, content_type)
        if status != 200:
            return jsonify(result), status
        words = ", ".join(result.get("text", []))[:2000]
        labels = ", ".join(str(x.get("name", "")) for x in result.get("labels", [])[:3])
        return jsonify(kind="image", summary=f"图片文字：{words or '未辨识到'}。粗分类：{labels or '未辨识到'}。",
                       note="本机 Vision OCR 与粗分类结果；尚非完整图片问答。")
    try:
        return jsonify(extract_document(file.filename or "", data))
    except (ValueError, zipfile.BadZipFile, ET.ParseError, ImportError) as exc:
        return jsonify(error=str(exc)), 400
    except Exception:
        return jsonify(error="无法抽取此文件的文字"), 422


@app.post("/api/chat")
def chat():
    guard = _require_login(api=True)
    if guard:
        return guard
    if not _same_origin():
        return jsonify(error="来源不符"), 403
    if not BRIDGE:
        return jsonify(error="此 Railway 预览尚未接入对话服务"), 503
    if request.mimetype != "application/json" or not request.content_length or request.content_length > MAX_JSON:
        return jsonify(error="文字格式不正确"), 415
    body = request.get_json(silent=True)
    if not isinstance(body, dict):
        return jsonify(error="文字格式不正确"), 400
    text = str(body.get("message", ""))[:4000]
    status, result = _bridge("/api/chat", json.dumps({"text": text, "history": []}).encode(), "application/json")
    return jsonify(result), status


@app.post("/api/transcribe")
def transcribe():
    guard = _require_login(api=True)
    if guard:
        return guard
    if not _same_origin():
        return jsonify(error="来源不符"), 403
    if not BRIDGE:
        return jsonify(error="此 Railway 预览尚未接入回合语音"), 503
    if request.mimetype != "audio/wav" or not request.content_length or not 44 <= request.content_length <= MAX_AUDIO:
        return jsonify(error="音讯格式不正确"), 415
    status, result = _bridge("/api/transcribe", request.get_data(), "audio/wav")
    return jsonify(result), status


@app.get("/")
@app.get("/<path:filename>")
def files(filename=""):
    if not filename:
        guard = _require_login()
        if guard:
            return guard
        return redirect("/chat", code=303)
    public = {"style.css", "login.js", "login-minimal.css", "assets/icon.svg", "assets/icon-180.png",
              "assets/icon-192.png", "assets/icon-512.png",
              "assets/dunni-icon-v1-180.png", "assets/dunni-icon-v1-192.png",
              "assets/dunni-icon-v1-512.png", "assets/dunni-login-loop.mp4",
              "assets/dunni-login-poster.jpg"}
    if filename not in public:
        guard = _require_login()
        if guard:
            return guard
    if filename in {"login.html", "server.py"} or filename.startswith("."):
        abort(404)
    return send_from_directory(ROOT, filename)


@app.errorhandler(413)
def too_large(_error):
    if request.path.startswith("/api/"):
        return jsonify(error="文件大小超过限制"), 413
    return "文件大小超过限制", 413


if __name__ == "__main__":
    if not _configured():
        raise SystemExit("Preview login variables are not configured")
    port = int(os.environ.get("PORT", "8080"))
    host = "127.0.0.1" if BRIDGE else "0.0.0.0"
    app.run(host=host, port=port, threaded=True)
